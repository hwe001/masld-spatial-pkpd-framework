from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT_DIR = ROOT / "outputs" / "supplementary_information"
OUT_DIR.mkdir(parents=True, exist_ok=True)
DIAGRAM = OUT_DIR / "supplementary_model_diagram.png"
DOCX = OUT_DIR / "Supplementary_Information.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def set_cell(cell, text, bold=False, size=8.2):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, headers, rows, size=8.2):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, bold=True, size=size)
        shade(t.rows[0].cells[i], "1F4E79")
        t.rows[0].cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value, size=size)
    for row in t.rows:
        trPr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        trPr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

def equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.name = "Cambria Math"
    r.font.size = Pt(10.5)
    r.italic = True

def make_diagram():
    fig, ax = plt.subplots(figsize=(12, 6.4), dpi=220)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6.4)
    ax.axis("off")
    colors = {"input": "#DCE6F1", "state": "#E2F0D9", "out": "#FCE4D6", "feedback": "#FFF2CC"}
    def box(x, y, w, h, label, kind):
        patch = FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.04,rounding_size=0.08",
                               linewidth=1.6, edgecolor="#1F4E79", facecolor=colors[kind])
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=10,
                fontweight="bold", wrap=True)
    def arrow(x1, y1, x2, y2, label=None, dashed=False):
        ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>", mutation_scale=16,
                                     linewidth=1.8, linestyle="--" if dashed else "-", color="#404040"))
        if label:
            ax.text((x1+x2)/2, (y1+y2)/2 + 0.16, label, ha="center", va="bottom",
                    fontsize=9, fontweight="bold", color="#404040")
    ax.add_patch(FancyBboxPatch((0.12, 2.48), 11.7, 3.45, boxstyle="round,pad=0.03,rounding_size=0.08",
                                linewidth=1.8, edgecolor="#1F4E79", facecolor="#EAF2F8", linestyle="--", zorder=0))
    ax.add_patch(FancyBboxPatch((0.35, 0.02), 10.1, 1.55, boxstyle="round,pad=0.03,rounding_size=0.08",
                                linewidth=1.8, edgecolor="#548235", facecolor="#EEF6E8", linestyle="--", zorder=0))
    ax.text(6, 6.15, "Supplementary model architecture", ha="center", fontsize=16, fontweight="bold")
    ax.text(9.0, 5.98, "NEW REDUCED PHARMACOLOGY ADAPTER", ha="center", va="center", fontsize=9,
            fontweight="bold", color="#1F4E79")
    ax.text(8.25, 1.50, "SEPARATE GADOXETATE CALIBRATION SCAFFOLD", ha="center", va="center", fontsize=9,
            fontweight="bold", color="#548235")
    ax.text(3.0, 5.65, "Ten-zone transport / obstruction prototype", ha="center", fontsize=12, fontweight="bold", color="#1F4E79")
    box(0.3, 4.65, 1.55, 0.75, "Input pulse\n(t <= 2 model units)", "input")
    box(2.25, 4.65, 1.7, 0.75, "Zonated\nuptake", "state")
    box(4.4, 4.65, 1.7, 0.75, "Parent pool\np_i", "state")
    box(6.55, 4.65, 1.9, 0.75, "Metabolism\n(CYP-like)", "state")
    box(8.9, 4.65, 1.7, 0.75, "Metabolite\npool m_i", "state")
    box(10.95, 4.65, 0.75, 0.75, "Bile", "out")
    arrow(1.85, 5.02, 2.25, 5.02)
    arrow(3.95, 5.02, 4.4, 5.02)
    arrow(6.1, 5.02, 6.55, 5.02)
    arrow(8.45, 5.02, 8.9, 5.02)
    arrow(10.6, 5.02, 10.95, 5.02, "export")
    box(2.4, 2.95, 2.0, 0.8, "Inherited bile-flow\nA/K/e velocity", "input")
    box(5.1, 2.95, 2.0, 0.8, "Injury state\nI_i", "feedback")
    box(7.8, 2.95, 2.0, 0.8, "Basolateral\nescape", "out")
    arrow(10.2, 4.65, 9.8, 3.75, "bile load")
    arrow(8.9, 4.65, 6.5, 3.75, "toxic load")
    arrow(7.1, 3.35, 7.8, 3.35, "ABCC3/4")
    arrow(5.1, 3.35, 4.4, 3.35, "flow feedback", dashed=True)
    arrow(4.4, 3.35, 4.4, 4.65, dashed=True)
    ax.text(3.0, 1.35, "Gadoxetate calibration scaffold", ha="center", fontsize=12, fontweight="bold", color="#1F4E79")
    box(0.6, 0.35, 1.55, 0.7, "Plasma\nP", "input")
    box(3.0, 0.35, 1.7, 0.7, "Hepatocyte\nH", "state")
    box(5.6, 0.35, 1.35, 0.7, "Bile\nB", "out")
    box(8.0, 0.35, 1.65, 0.7, "Renal\nclearance", "out")
    arrow(2.15, 0.7, 3.0, 0.7, "uptake")
    arrow(4.7, 0.7, 5.6, 0.7, "bile")
    arrow(3.0, 0.48, 2.15, 0.48, "back")
    ax.add_patch(FancyArrowPatch((1.35, 0.18), (8.0, 0.18), arrowstyle="-|>", mutation_scale=16,
                                 linewidth=1.8, color="#404040"))
    ax.text(6.7, 0.05, "parallel renal route", ha="center", va="top", fontsize=9, fontweight="bold", color="#404040")
    fig.savefig(DIAGRAM, bbox_inches="tight", facecolor="white")
    plt.close(fig)

make_diagram()
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)
styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].paragraph_format.space_after = Pt(5)
styles["Normal"].paragraph_format.line_spacing = 1.05
for name in ("Title", "Heading 1", "Heading 2"):
    styles[name].font.name = "Arial"
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(19)
styles["Heading 1"].font.size = Pt(13)
styles["Heading 2"].font.size = Pt(10.5)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Supplementary Information")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Testing the pharmacological consequences of liver zonation in MASLD using a spatial PK/PD framework")
r.italic = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Harvey Ho | Rātā BioSim Lab, Auckland, New Zealand | harvey@ratalab.nz")
r.bold = True

doc.add_heading("S1. Purpose and scope", level=1)
doc.add_paragraph(
    "This Supplementary Information records the model structure, state definitions, parameter values, scenario design, and code organization used for the manuscript. The implementation contains two related but distinct components. The coupled transport prototype represents a ten-position CV-to-PV reduced mesh with zonated transporter fields, metabolite handling, bile advection, obstruction, and an injury state. The gadoxetate scaffold is a separate three-compartment calibration model with plasma, hepatocyte, and bile amounts plus parallel renal clearance. The normalized transport prototype is not a clinical dosing model and the coefficient fields are hypothesis parameters pending functional calibration."
)

doc.add_heading("S2. Model architecture", level=1)
doc.add_picture(str(DIAGRAM), width=Inches(6.9))
p = doc.paragraphs[-1]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph("Figure S1. Computational architecture of the coupled transport prototype and the separate gadoxetate calibration scaffold. Disease state and transporter fields modify uptake, export, metabolism, injury susceptibility, and basolateral escape; obstruction scales local bile velocity. The upper layer is a new reduced pharmacology adapter using Michael Kuecken's input profile, not a reproduction of the original full bile-flow solver.")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.runs[0].italic = True

doc.add_heading("S3. Coupled transport prototype", level=1)
doc.add_paragraph(
    "The prototype uses ten reduced positions ordered from central to portal: central, central, mid, mid, mid, periportal, periportal, portal, portal, portal. Michael Kuecken's normal-control and steatosis input profiles provide the local quantities A, K, epithelial volume fraction e, and connectivity. A normalized carrier velocity is calculated as v_i = (A_i/K_i)e_i^-1 c_i / median[(A/K)e^-1 c]. Obstruction scales this velocity by (1 - obstruction)."
)
doc.add_paragraph(
    "The implemented state vector is y = [p_1...p_10, m_1...m_10, b_1...b_10, I_1...I_10], where p is the local parent pool, m is the metabolite/active hepatocyte pool, b is bile-side metabolite amount, and I is a normalized injury state. In this prototype, p_i is not a full plasma compartment; it is a zone-local parent pool driven by the input pulse."
)
doc.add_heading("S3.1 Governing equations", level=2)
for eq in [
    "u_i = 0.10 * U_i * q_u * exp(-1.7 i/n) * 1(t <= 2)",
    "r_i = 0.40 * CYP_i * p_i/(0.35 + p_i)",
    "x_i = f_m * C_i * X_i * 0.30 * m_i/(0.25 + m_i)",
    "dp_i/dt = u_i - r_i - 0.025 p_i",
    "dm_i/dt = r_i - x_i - 0.020 m_i",
    "db_i/dt = x_i/e_i + v_(i-1)b_(i-1) - v_i b_i",
    "dI_i/dt = 0.020 J_i L_i - 0.004 I_i + 0.003 E_i I_i",
]:
    equation(doc, eq)
doc.add_paragraph(
    "Here U_i is the aggregated uptake capacity, C_i is the aggregated canalicular capacity, E_i is the aggregated basolateral escape capacity, J_i is the disease injury factor, f_m is the MRP2 activity factor, q_u is the uptake inhibition factor, and X_i is the disease export factor. The toxic load is L_i = max(0, m_i - 0.05 GSH_i) + 0.003 b_i. Local flow is max(0.08, v_i(1 - 0.80 I_i)). The first bile position receives an outflow term -v_1 b_1; upstream-to-downstream indexing follows the implementation's reduced CV-to-PV ordering."
)

doc.add_heading("S3.2 Fixed numerical settings", level=2)
add_table(doc, ["Item", "Value", "Interpretation"], [
    ("Time interval", "0-60 model units", "Transport prototype simulation window"),
    ("Time grid", "1201 points; delta t = 0.05", "Output grid supplied to solve_ivp"),
    ("Input pulse", "1.0 for t <= 2.0; 0 otherwise", "Dimensionless normalized input"),
    ("Integrator", "scipy.integrate.solve_ivp", "Adaptive ODE integration"),
    ("Relative tolerance", "1 x 10^-7", "Solver setting"),
    ("Absolute tolerance", "1 x 10^-9", "Solver setting"),
    ("Uptake inhibition", "q_u = 0.5", "50% activity condition"),
    ("MRP2 activity", "f_m = 1.0 or 0.35", "Baseline or 65% inhibition"),
    ("Obstruction", "0.0, 0.5; LHS 0-0.75", "Fractional velocity reduction"),
])

doc.add_heading("S4. Spatial transporter and disease fields", level=1)
doc.add_paragraph(
    "The four-zone fields are dimensionless relative capacities. They are explicit hypothesis parameters, not direct compound-specific functional measurements. The ten-position mesh expands these values using the zone labels listed above. Uptake is the mean of SLC10A1, SLCO1B1, and SLCO1B3; canalicular export is the mean of ABCB11, ABCC2, and ABCB4; basolateral escape is the mean of ABCC3 and ABCC4; ductular capacity is the mean of CFTR and SLC4A2."
)
field_rows = [
    ("Control", "Uptake", "1.00", "1.00", "0.98", "0.95"),
    ("Control", "Canalicular", "1.00", "1.00", "1.00", "1.00"),
    ("Control", "Escape", "1.00", "1.00", "1.00", "1.00"),
    ("MASL", "Uptake", "0.98", "0.96", "0.92", "0.88"),
    ("MASL", "Canalicular", "1.00", "0.97", "0.92", "0.88"),
    ("MASL", "Escape", "1.05", "1.08", "1.12", "1.16"),
    ("MASH", "Uptake", "0.95", "0.90", "0.82", "0.74"),
    ("MASH", "Canalicular", "1.00", "0.90", "0.78", "0.68"),
    ("MASH", "Escape", "1.10", "1.20", "1.35", "1.50"),
]
add_table(doc, ["State", "Field", "Portal", "Periportal", "Mid", "Central"], field_rows)
doc.add_paragraph("The individual gene-level fields, including ABCB11, ABCC2, ABCB4, CFTR, and SLC4A2, are preserved in zonated_transport_genes.py. The control, MASL, and MASH metabolism factors are respectively [1.00, 1.05, 1.00, 0.95], [0.95, 0.92, 0.88, 0.82], and [0.90, 0.85, 0.78, 0.68] from portal to central. Injury factors are [1.00, 1.00, 1.00, 1.00], [1.00, 1.05, 1.10, 1.15], and [1.05, 1.15, 1.30, 1.45].")

doc.add_heading("S5. Michael input profiles", level=1)
doc.add_paragraph(
    "The full ten-row input arrays are embedded in couple_michael_bileflow.py. A, K, epithelial volume fraction e, and connectivity c are retained as source inputs for the normalized carrier field. The source file is the authoritative machine-readable parameter record; the summary below identifies the implementation and its use."
)
add_table(doc, ["Input", "Normal control", "Steatosis", "Use"], [
    ("Rows", "10 CV-to-PV positions", "10 CV-to-PV positions", "Reduced spatial carrier"),
    ("A", "Source array in code", "Source array in code", "Conductance numerator"),
    ("K", "Source array in code", "Source array in code", "Conductance denominator"),
    ("e", "0.0310-0.0419", "0.00287-0.0411", "Velocity normalization"),
    ("Connectivity c", "0.7719-0.9824", "0.0100-0.9436", "Velocity weighting"),
    ("Velocity scaling", "median-normalized", "median-normalized", "Additional obstruction multiplier"),
], size=8.0)
from couple_michael_bileflow import INPUTS
for profile_name, profile_label in (("normal control", "Normal-control profile"), ("steatosis", "Steatosis profile")):
    doc.add_heading(profile_label + " exact values", level=2)
    exact_rows = []
    for i, row in enumerate(INPUTS[profile_name], start=1):
        exact_rows.append((i, f"{row[0]:.12g}", f"{row[1]:.12g}", f"{row[2]:.8f}", f"{row[3]:.6f}"))
    add_table(doc, ["Position", "A", "K", "e", "Connectivity"], exact_rows, size=7.4)

doc.add_heading("S6. Gadoxetate calibration scaffold", level=1)
doc.add_paragraph(
    "The gadoxetate model is a separate three-state ODE system with a 24-hour simulation window. Amounts are in mmol for a 70 kg adult. The recommended dose is represented as 0.025 mmol/kg x 70 kg = 1.75 mmol. The label-inspired volume and clearance quantities are converted to an hour-based rate scale."
)
for eq in [
    "dP/dt = -(CL_renal/V_P)*renal*P - uptake*P + back*H",
    "dH/dt = uptake*P - (back + bile)*H",
    "dB/dt = bile*H",
]:
    equation(doc, eq)
add_table(doc, ["Scenario", "uptake", "back", "bile", "renal factor", "Purpose"], [
    ("Healthy", "1.10", "0.12", "0.112", "1.00", "Approximate 50:50 renal/hepatobiliary recovery"),
    ("MASL/MASH exploratory", "0.88", "0.16", "0.090", "1.00", "Reduced uptake and biliary handling"),
    ("Severe hepatic impairment", "0.35", "0.24", "0.025", "1.30", "Approximate 5% hepatobiliary recovery"),
])
add_table(doc, ["Fixed quantity", "Value"], [
    ("Body mass", "70 kg"),
    ("Dose", "1.75 mmol"),
    ("Vd approximation", "0.21 L/kg x 70 kg = 14.7 L"),
    ("Total clearance", "0.250 L/min x 60 = 15.0 L/h"),
    ("Renal clearance", "0.120 L/min x 60 = 7.2 L/h"),
    ("Hepatic remainder", "7.8 L/h"),
    ("Time grid", "0-24 h; 2401 points"),
])

doc.add_heading("S7. Analysis design and endpoints", level=1)
doc.add_paragraph("The counterfactual comparison holds the mean spatial capacity constant and compares zonated versus well-mixed fields. Primary quantities are bile AUC, hepatocyte parent AUC, bile Cmax, time to bile Cmax, peak injury, and terminal injury.")
doc.add_paragraph("A reproducible transporter stress test compares each disease state with its own baseline. At 50% uptake activity, bile-AUC ratios are 0.501, 0.502, and 0.502 for control, MASL, and MASH, while local parent-pool AUC ratios are 0.480, 0.480, and 0.482. At MRP2 activity of 0.35, bile-AUC ratios are 0.964, 0.939, and 0.911, while local parent-pool AUC ratios remain 1.000. These are normalized model outputs; they are not plasma AUC ratios or clinical DDI predictions because the current prototype has no full plasma victim-drug compartment.")
add_table(doc, ["Analysis", "Scenarios / ranges", "Output"], [
    ("Counterfactual", "Zonated vs well mixed; baseline, 50% uptake inhibition, 65% MRP2 inhibition", "Relative bile and hepatocyte AUC differences"),
    ("Scenario grid", "MRP2 = 1.0 or 0.35; obstruction = 0 or 0.5", "Bile Cmax, AUC, tmax, peak and terminal injury"),
    ("LHS sensitivity", "160 samples per normal-control/steatosis profile; MRP2 0.20-1.20; obstruction 0-0.75", "Exposure and injury distributions"),
    ("Gadoxetate", "Healthy, exploratory MASL/MASH, severe impairment", "24-hour route recovery and amount-time curves"),
])

doc.add_heading("S8. Code map and reproducibility", level=1)
add_table(doc, ["File", "Role"], [
    ("couple_michael_bileflow.py", "Defines input profiles, ten-position mapping, disease factors, ODE, obstruction, export, bile advection, and injury feedback."),
    ("zonated_transport_genes.py", "Stores four-zone gene-level relative fields and aggregates uptake, canalicular, escape, and ductular capacities."),
    ("compare_zonated_wellmixed.py", "Runs the mean-preserving spatial counterfactual and writes main-manuscript Figure 2 plus metric CSV."),
    ("run_jpkpd_analysis.py", "Runs scenario grids and seeded Latin-hypercube sensitivity analysis; writes main-manuscript Figure 3 and CSV outputs."),
    ("simulate_gadoxetate_pk.py", "Runs the separate plasma-hepatocyte-bile calibration scaffold and writes main-manuscript Figure 4 and CSV outputs."),
    ("render_ddi_efficacy_figure.py", "Plots the transporter stress-test ratios and the illustrative efficacy-mapping extension used in main-manuscript Figure 5."),
    ("create_second_paper_draft.py", "Builds the main manuscript DOCX and inserts the five main figures."),
    ("create_supplementary_information.py", "Builds this supplementary document and its model architecture diagram."),
])
doc.add_paragraph("From the project root, run compare_zonated_wellmixed.py, run_jpkpd_analysis.py, simulate_gadoxetate_pk.py, and render_ddi_efficacy_figure.py to regenerate numerical outputs and figures. Then run create_second_paper_draft.py and create_supplementary_information.py to regenerate the DOCX files. The sensitivity generator uses numpy.random.default_rng(20260927).")
doc.add_paragraph("The model is intentionally transparent rather than clinically validated. The next calibration release should replace relative transporter fields with zone-resolved protein or functional measurements, retain uncertainty distributions, and test predictions against independent gadoxetate and transporter-probe data.")

doc.add_heading("S9. Relationship to the main manuscript", level=1)
doc.add_paragraph("Figure S1 and the parameter tables provide implementation detail for the Methods and Results sections. The main manuscript reports normalized outputs and interprets them as structural-sensitivity results. No supplementary parameter should be interpreted as a clinical concentration, patient-specific coefficient, or validated toxicity threshold.")

doc.save(DOCX)
print(DOCX)
